"""Read-only native source adapters. Object IDs survive Agent round-trips.
SVG outline comments are hints, PPTX group transforms are explicit, PDF text is native.
"""
from __future__ import annotations
import io,math,re,zipfile
from pathlib import Path
from lxml import etree as ET
from PIL import Image
import pymupdf as fitz
from .common import FigureError, digest_file, digest_json
from .assets import svg_root,NS,XLINK,svg_pdf_bytes

LIMITS={'bytes':80_000_000,'pages':40,'objects':12000,'zip_uncompressed':200_000_000}


def matmul(a,b):
    return [sum(a[i*3+k]*b[k*3+j] for k in range(3)) for i in range(3) for j in range(3)]


def point(m,x,y):return [m[0]*x+m[1]*y+m[2],m[3]*x+m[4]*y+m[5]]

def bbox_transform(m,r):
    p=[point(m,x,y) for x in (r[0],r[2]) for y in (r[1],r[3])]
    return [min(q[0] for q in p),min(q[1] for q in p),max(q[0] for q in p),max(q[1] for q in p)]

IDENTITY=[1,0,0,0,1,0,0,0,1]


def _rotation(cx,cy,deg,fh=False,fv=False):
    rad=math.radians(deg);c,s=math.cos(rad),math.sin(rad);sx=-1 if fh else 1;sy=-1 if fv else 1
    return matmul([1,0,cx,0,1,cy,0,0,1],matmul([c,-s,0,s,c,0,0,0,1],matmul([sx,0,0,0,sy,0,0,0,1],[1,0,-cx,0,1,-cy,0,0,1])))


def group_matrix(xfrm):
    """OOXML child EMU -> parent EMU, including nested scale/flip/rotation."""
    def vals(name,attrs):
        el=xfrm.find('{http://schemas.openxmlformats.org/drawingml/2006/main}'+name)
        if el is None:raise FigureError('pptx_group_transform','组变换缺少'+name)
        return [float(el.get(a)) for a in attrs]
    x,y=vals('off',['x','y']);w,h=vals('ext',['cx','cy']);cx,cy=vals('chOff',['x','y']);cw,ch=vals('chExt',['cx','cy'])
    if cw<=0 or ch<=0:raise FigureError('pptx_group_transform','组的子坐标范围无效')
    base=[w/cw,0,x-cx*w/cw,0,h/ch,y-cy*h/ch,0,0,1]
    r=_rotation(x+w/2,y+h/2,float(xfrm.get('rot','0'))/60000,xfrm.get('flipH')=='1',xfrm.get('flipV')=='1')
    return matmul(r,base)


def _pdf_value(v):
    if isinstance(v,(str,int,float,bool,type(None))):return v
    if isinstance(v,dict):return {str(k):_pdf_value(x) for k,x in v.items() if not isinstance(x,bytes)}
    if isinstance(v,(fitz.Rect,fitz.Point,fitz.Quad,fitz.Matrix,tuple,list)):return [_pdf_value(x) for x in v]
    return str(v)


def pdf_pages(path,folder):
    out=[]
    with fitz.open(path) as doc:
        if doc.needs_pass:raise FigureError('encrypted_pdf','加密PDF需先提供可读副本')
        if len(doc)>LIMITS['pages']:raise FigureError('page_limit','超过一次读取的页数上限')
        for i,p in enumerate(doc):
            texts=[]
            for bi,block in enumerate(p.get_text('dict',flags=fitz.TEXT_PRESERVE_WHITESPACE|fitz.TEXT_PRESERVE_LIGATURES)['blocks']):
                if block.get('type')!=0:continue
                for li,line in enumerate(block.get('lines',[])):
                    spans=line.get('spans',[])
                    texts.append({'id':f'p{i+1}-text-{bi}-{li}','kind':'native-text','text':''.join(x['text'] for x in spans),
                        'rect':list(line['bbox']),'direction':list(line['dir']),
                        'runs':[{'text':x['text'],'font':x['font'],'size_pt':x['size'],'flags':x['flags'],'rect':list(x['bbox'])} for x in spans]})
            drawings=p.get_drawings();objs=[]
            if len(drawings)+len(texts)>LIMITS['objects']:raise FigureError('object_limit','PDF对象数量超过当前批次限额')
            for j,row in enumerate(drawings):
                objs.append({'id':f'p{i+1}-path-{j}','kind':'native-path','rect':list(row['rect']),
                             'geometry':_pdf_value(row),'shape_hint':'rectangle' if any(t[0]=='re' for t in row['items']) else None})
            image_rows=[{'id':f'p{i+1}-image-{j}','kind':'embedded-image','rect':list(v['bbox']),
                'native_size':[v['width'],v['height']],'transform':list(v['transform'])} for j,v in enumerate(p.get_image_info())]
            preview=f'page-{i+1}.png';scale=min(2.,1600/max(p.rect.width,p.rect.height))
            p.get_pixmap(matrix=fitz.Matrix(scale,scale),alpha=False).save(folder/preview)
            # Rotated pages keep explicit coordinate transform; Agent can't silently use display pixels as PDF points.
            out.append({'id':f'page-{i+1}','kind':'pdf','size':[p.cropbox.width,p.cropbox.height],'display_size':[p.rect.width,p.rect.height],
                'coordinate_unit':'pt-unrotated','page_index':i,'rotation':p.rotation,
                'rotation_matrix':list(p.rotation_matrix),'cropbox':list(p.cropbox),
                'preview':preview,'preview_scale':scale,'objects':texts+objs+image_rows,'chart_candidates':[],
                'warnings':[{'code':'pdf_chart_binding_pending','message':'已提取原生对象；PDF统计图自动去旧字/图例仍需显式安全适配，不自动栅格化'}],
                'counts':{'texts':len(texts),'paths':len(objs),'images':len(image_rows)}})
    return out


def pptx_pages(path,folder):
    from pptx import Presentation
    with zipfile.ZipFile(path) as z:
        if sum(n.file_size for n in z.infolist())>LIMITS['zip_uncompressed']:raise FigureError('archive_size_limit','PPTX解压大小超限')
    prs=Presentation(path)
    if len(prs.slides)>LIMITS['pages']:raise FigureError('page_limit','幻灯片数量超限')
    out=[];a='{http://schemas.openxmlformats.org/drawingml/2006/main}'
    for si,slide in enumerate(prs.slides):
        objects=[];warnings=[]
        def walk(shapes,parent=None,m=IDENTITY):
            for sh in shapes:
                ident=f's{si+1}-shape-{sh.shape_id}';r=[float(sh.left),float(sh.top),float(sh.left+sh.width),float(sh.top+sh.height)]
                xfrm=sh.element.find('.//'+a+'xfrm')
                rotation=float(sh.rotation);fh=xfrm is not None and xfrm.get('flipH')=='1';fv=xfrm is not None and xfrm.get('flipV')=='1'
                localrot=_rotation((r[0]+r[2])/2,(r[1]+r[3])/2,rotation,fh,fv)
                world=matmul(m,localrot);rect=[q/12700 for q in bbox_transform(world,r)]
                row={'id':ident,'kind':'native-shape','native_id':sh.shape_id,'name':sh.name,'parent_id':parent,
                    'rect':rect,'local_rect_emu':r,'world_matrix':world,'rotation':rotation,'shape_type':str(sh.shape_type),
                    'text':sh.text if sh.has_text_frame else None}
                geom=sh.element.find('.//'+a+'prstGeom');row['shape_hint']=geom.get('prst') if geom is not None else 'custom-or-group'
                if sh.has_text_frame:
                    row['runs']=[{'text':run.text,'bold':run.font.bold,'italic':run.font.italic,
                       'font':run.font.name,'size_pt':run.font.size.pt if run.font.size else None}
                       for para in sh.text_frame.paragraphs for run in para.runs]
                if hasattr(sh,'begin_x'):
                    row['kind']='native-connector';row['endpoints']=[[v/12700 for v in point(world,r[0],r[1])],[v/12700 for v in point(world,r[2],r[3])]]
                    for which,tag in [('from','stCxn'),('to','endCxn')]:
                        q=sh.element.find('.//'+a+tag)
                        row[which+'_ref']=f's{si+1}-shape-'+q.get('id') if q is not None else None
                    row['head']=_end(sh.element,a+'headEnd');row['tail']=_end(sh.element,a+'tailEnd')
                    row['topology_status']='native-bound' if row['from_ref'] and row['to_ref'] else 'unbound-endpoints-needs-agent'
                if sh.shape_type==13:
                    row['kind']='embedded-image';row['native_size']=list(sh.image.size)
                    row['crop']=[sh.crop_left,sh.crop_top,sh.crop_right,sh.crop_bottom]
                    media=f'{ident}.{sh.image.ext}';(folder/media).write_bytes(sh.image.blob);row['media']=media
                    row['sha256']=digest_file(folder/media)
                objects.append(row)
                if len(objects)>LIMITS['objects']:raise FigureError('object_limit','PPTX对象数量超限')
                if hasattr(sh,'shapes'):
                    gm=group_matrix(xfrm) if xfrm is not None else IDENTITY
                    walk(sh.shapes,ident,matmul(m,gm))
        walk(slide.shapes)
        out.append({'id':f'slide-{si+1}','kind':'pptx','size':[prs.slide_width/12700,prs.slide_height/12700],
            'coordinate_unit':'pt','objects':objects,'chart_candidates':[],'preview':None,
            'warnings':[{'code':'pptx_not_rendered','message':'原生对象证据非完整PPT渲染；已有原PDF/截图用于图意核对，不能用几何缩略图冒充原稿'}],
            'counts':{'texts':sum(bool(q.get('text')) for q in objects),'objects':len(objects),'connectors':sum(q['kind']=='native-connector' for q in objects)}})
    return out


def _end(el,tag):
    q=el.find('.//'+tag);return q.get('type','none') if q is not None else 'none'


def simple_path_points(data):
    """Only M/L, absolute or relative. Curves are NOT approximated as data points."""
    tokens=re.findall(r'[A-Za-z]|[-+]?(?:\d*\.\d+|\d+\.?)(?:[eE][-+]?\d+)?',data or '')
    if any(t.isalpha() and t not in 'MmLlZz' for t in tokens):return []
    pts=[];i=0;cmd='';x=y=0.
    try:
        while i<len(tokens):
            if tokens[i].isalpha():cmd=tokens[i];i+=1
            if cmd in 'Zz':break
            if not cmd or i+1>=len(tokens):return []
            xx,yy=float(tokens[i]),float(tokens[i+1]);i+=2
            if cmd.islower():x+=xx;y+=yy
            else:x,y=xx,yy
            pts.append([x,y])
    except (ValueError,IndexError):return []
    return pts


def _comment_text(el):
    hints=[q.text.strip() for q in el if not isinstance(q.tag,str) and q.text and q.text.strip()]
    native=[''.join(q.itertext()).strip() for q in el.findall('.//s:text',NS)]
    return ''.join(native) if native else ' '.join(hints)


def _group_bounds(el,root,pad=1.):
    # Supports ordinary Matplotlib line / marker groups; excludes defs and font outlines.
    pts=[]
    for q in el.iter():
        if not isinstance(q.tag,str) or any(isinstance(a.tag,str) and ET.QName(a).localname=='defs' for a in q.iterancestors()):continue
        tag=ET.QName(q).localname
        if tag=='path':pts+=simple_path_points(q.get('d'))
        if tag=='use':
            x,y=float(q.get('x','0')),float(q.get('y','0'));pts.extend([[x-4,y-4],[x+4,y+4]])
    if not pts:return None
    return [min(p[0] for p in pts)-pad,min(p[1] for p in pts)-pad,max(p[0] for p in pts)+pad,max(p[1] for p in pts)+pad]


def svg_pages(path,folder):
    root=svg_root(path);view=[float(x) for x in root.get('viewBox','').replace(',',' ').split()]
    if len(view)!=4:
        # Inventory only: no inference that an anonymous outlined SVG is a chart.
        from .source_outline_chart import primitive_inventory
        inventory=primitive_inventory(path)
        return [{'id':'svg-1','kind':'svg','size':inventory['viewbox'][2:],'viewbox':inventory['viewbox'],
                 'coordinate_unit':'svg-user-unit','preview':None,'objects':inventory['objects'],
                 'outline_inventory':inventory,'chart_candidates':[],
                 'warnings':[{'code':'outlined_svg_registration_required','message':'仅提供原图元证据；需显式完整登记并核对转曲原文，不自动识图'}],
                 'counts':{'objects':len(inventory['objects']),'charts':0}}]
    ids=[q.get('id') for q in root.iter() if isinstance(q.tag,str) and q.get('id')]
    if len(ids)!=len(set(ids)):raise FigureError('svg_duplicate_ids','SVG ID重复，不能安全绑定')
    objects=[]
    for q in root.findall('.//s:g',NS):
        ident=q.get('id','')
        if ident.startswith('text_'):
            text=_comment_text(q);anc=[a.get('id') for a in q.iterancestors() if a.get('id')]
            objects.append({'id':ident,'kind':'svg-text-hint','text':text,'source_group':ident,'ancestors':anc,
                'confidence':'native-text' if q.findall('.//s:text',NS) else 'outlined-comment-needs-visual-confirmation',
                'transform':[e.get('transform') for e in q.iter() if isinstance(e.tag,str) and e.get('transform')][:4]})
    charts=[]
    for axes in root.findall('.//s:g',NS):
        aid=axes.get('id','')
        if not aid.startswith('axes_'):continue
        # Do not silently apply unknown inherited transforms to original coordinates.
        trans=[q.get('transform') for q in [axes,*axes.iterancestors()] if q.get('transform')]
        if trans:continue
        patch=next((q for q in axes.findall('s:g',NS) if q.get('id','').startswith('patch_')),None)
        pts=simple_path_points(patch[0].get('d')) if patch is not None and len(patch) else []
        if len(pts)<4:continue
        frame=[min(x for x,y in pts),min(y for x,y in pts),max(x for x,y in pts),max(y for x,y in pts)]
        ticks={'x':[],'y':[]};titles={};tick_lengths=[]
        for axgroup in axes.findall('s:g',NS):
            if not axgroup.get('id','').startswith('matplotlib.axis_'):continue
            orientation=None
            for tq in axgroup.findall('s:g',NS):
                tid=tq.get('id','')
                if not tid.startswith(('xtick_','ytick_')):continue
                orientation=tid[0]
                uses=[u for u in tq.findall('.//s:use',NS) if u.get('x') is not None and u.get('y') is not None]
                if not uses:continue
                u=uses[0];position=float(u.get(orientation));tg=next((g for g in tq.findall('s:g',NS) if g.get('id','').startswith('text_')),None)
                if tg is None:continue
                ticks[orientation].append({'id':tid,'position':position,'text_ref':tg.get('id'),'text_hint':_comment_text(tg)})
                target=root.xpath('//*[@id=$id]',id=u.get(XLINK,u.get('href',''))[1:])
                if target:
                    tp=simple_path_points(target[0].get('d'))
                    if len(tp)==2:tick_lengths.append(math.dist(tp[0],tp[1]))
            tg=next((g for g in axgroup.findall('s:g',NS) if g.get('id','').startswith('text_')),None)
            if orientation and tg is not None:titles[orientation]={'ref':tg.get('id'),'text_hint':_comment_text(tg)}
        legend=next((q for q in axes.findall('s:g',NS) if q.get('id','').startswith('legend_')),None)
        series=[q for q in axes.findall('s:g',NS) if q.get('id','').startswith(('line2d_','PathCollection_'))]
        rows=[]
        if legend is not None:
            children=list(legend)
            for idx,q in enumerate(children):
                if not isinstance(q.tag,str) or not q.get('id','').startswith('text_'):continue
                sw=next((v for v in reversed(children[:idx]) if isinstance(v.tag,str) and v.get('id','').startswith(('line2d_','PathCollection_'))),None)
                if sw is None:continue
                bounds=_group_bounds(sw,root)
                if bounds:rows.append({'id':legend.get('id')+'-row-'+str(len(rows)+1),'label_ref':q.get('id'),
                    'text_hint':_comment_text(q),'swatch_group':sw.get('id'),'swatch_clip':bounds})
        protected=[]
        for q in series:
            for p in q.findall('.//s:path',NS):
                if any(ET.QName(a).localname=='defs' for a in p.iterancestors()):continue
                pts=simple_path_points(p.get('d'))
                if pts:protected.append(pts)
        charts.append({'id':aid,'kind':'matplotlib-svg-chart','axes_id':aid,'frame':frame,'plot_clip':frame,
            'x_ticks':ticks['x'],'y_ticks':ticks['y'],'titles':titles,'legend_id':legend.get('id') if legend is not None else None,
            'legend_rows':rows,'series_groups':[q.get('id') for q in series],'protected_path_points':protected,
            'tick_length':max(tick_lengths) if tick_lengths else None,'confidence':'native-geometry-semantic-confirmation-required'})
    outline_inventory=None
    if not charts:
        from .source_outline_chart import primitive_inventory
        try:outline_inventory=primitive_inventory(path)
        except FigureError:pass  # Existing unsupported native SVG behavior stays unchanged.
    # Render only once for visual evidence, not N isolated conversions for N labels.
    raw=svg_pdf_bytes(root)
    with fitz.open(stream=raw,filetype='pdf') as doc:
        p=doc[0];p.get_pixmap(matrix=fitz.Matrix(min(2.,1600/p.rect.width),min(2.,1600/p.rect.width)),alpha=False).save(folder/'svg-preview.png')
    return [{'id':'svg-1','kind':'svg','size':view[2:],'viewbox':view,'coordinate_unit':'svg-user-unit','preview':'svg-preview.png',
        'objects':objects,**({'outline_inventory':outline_inventory} if outline_inventory else {}),'chart_candidates':charts,'warnings':[] if charts else [{'code':'svg_structure_unrecognized','message':'已保留原文/对象，未假装任意SVG都是Matplotlib图'}],
        'counts':{'text_hints':len(objects),'charts':len(charts)},'analysis':{'svg_full_preview_conversions':1,'ocr_calls':0}}]


def extract_native(path:Path,folder:Path):
    if path.stat().st_size>LIMITS['bytes']:raise FigureError('source_size_limit','源文件超过本候选输入上限')
    suffix=path.suffix.lower()
    if suffix=='.pdf':return pdf_pages(path,folder)
    if suffix=='.pptx':return pptx_pages(path,folder)
    if suffix=='.svg':return svg_pages(path,folder)
    raise FigureError('unsupported_native_source','原生对象适配器不支持此类型',extension=suffix)
