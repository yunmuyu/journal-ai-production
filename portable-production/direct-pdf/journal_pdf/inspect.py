"""Read-only inventory. PPTX shape extraction is NOT a full fidelity renderer."""
from pathlib import Path
from PIL import Image
import pymupdf as fitz
from pptx import Presentation
from .common import digest_file,FigureError
from .assets import svg_root,NS


def inspect_file(path:Path):
    result={'file':path.name,'sha256':digest_file(path),'bytes':path.stat().st_size,'read_only':True}
    ext=path.suffix.lower()
    if ext=='.pdf':
        d=fitz.open(path);pages=[]
        for i,p in enumerate(d):
            pages.append({'page':i+1,'size_pt':list(p.rect),'text':p.get_text(),
                          'fonts':[{'name':r[3],'type':r[2]} for r in p.get_fonts()],
                          'image_resources':len(p.get_images()),'vector_paths':len(p.get_drawings())})
        d.close();return result|{'kind':'pdf','pages':pages}
    if ext=='.pptx':
        prs=Presentation(path)
        def walk(shapes,parent=None):
            out=[]
            for s in shapes:
                row={'id':s.shape_id,'name':s.name,'shape_type':str(s.shape_type),'parent_id':parent,
                     'local_rect_pt':[s.left/12700,s.top/12700,(s.left+s.width)/12700,(s.top+s.height)/12700],
                     'rotation':s.rotation,'text':s.text if s.has_text_frame else None,
                     'warning':'group child coordinates are LOCAL; do not draw as page coordinates' if parent else None}
                if hasattr(s,'shapes'):row['children']=walk(s.shapes,s.shape_id)
                if s.shape_type==9:
                    row['connector_xml']=s.element.xml
                out.append(row)
            return out
        return result|{'kind':'pptx','size_pt':[prs.slide_width/12700,prs.slide_height/12700],
                       'slides':[{'slide':i+1,'shapes':walk(s.shapes)} for i,s in enumerate(prs.slides)],
                       'full_fidelity_conversion':False,'requires_topology_review':True}
    if ext=='.svg':
        root=svg_root(path);texts=[]
        for e in root.findall('.//s:g',NS):
            if e.get('id','').startswith('text_'):
                comments=[c.text.strip() for c in e if not isinstance(c.tag,str)]
                texts.append({'id':e.get('id'),'comments':comments,'not_authoritative_without_visual_check':True})
        return result|{'kind':'svg','viewbox':root.get('viewBox'),'outlined_text_groups':texts,
                       'groups':[e.get('id') for e in root.findall('.//s:g',NS) if e.get('id')]}
    if ext in ('.png','.jpg','.jpeg','.tif','.tiff','.bmp'):
        with Image.open(path) as im:return result|{'kind':'raster','size_px':list(im.size),'mode':im.mode}
    raise FigureError('unsupported_inspection','该文件类型尚未支持读取',extension=ext)
