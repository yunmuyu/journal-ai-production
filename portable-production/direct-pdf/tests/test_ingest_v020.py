import copy,json,math
from pathlib import Path
import pytest
import pymupdf as fitz
from PIL import Image,ImageDraw,ImageFont
from pptx import Presentation
from pptx.util import Inches
from lxml import etree as ET
from journal_pdf.common import FigureError,read_json,digest_file,atomic_json
from journal_pdf.ingest_native import group_matrix,point,simple_path_points,extract_native
from journal_pdf.ingest_geometry import ranges,raster_candidates
from journal_pdf.agent_spec import compile_reply_figure,PROTOCOL
from journal_pdf.agent_entry import prepare,complete,review,run_with_agent
from journal_pdf.graph_layout import graph
from journal_pdf.scene import audit_scene
from journal_pdf.assets import Assets
from journal_pdf.render import render_one


def source_pdf(tmp_path):
    src=tmp_path/'source.pdf';d=fitz.open();p=d.new_page(width=220,height=180)
    for y,text in [(30,'Input'),(85,'Check'),(140,'Output')]:
        p.draw_rect([80,y-13,140,y+8],width=.5);p.insert_text((87,y),text,fontsize=10)
    d.save(src);d.close();return src


def reply_for_job(job,*,author='Synthetic',id='native-flow'):
    req=read_json(job/'agent-request.json');s=req['sources'][0];page=s['pages'][0];texts=page['text_objects']
    return {'protocol':PROTOCOL,'job_id':req['job_id'],'request_sha256':req['request_sha256'],'figures':[
        {'id':id,'author':author,'source_id':s['source_id'],'page_id':page['id'],'kind':'diagram',
        'review':{'source_checked':True,'labels_checked':True,'topology_checked':True},
        'nodes':[{'id':'n'+str(i),'ref':t['id'],'shape':'rectangle'} for i,t in enumerate(texts)],
        'edges':[{'id':'e'+str(i),'from':'n'+str(i),'to':'n'+str(i+1)} for i in range(len(texts)-1)]}], 'exceptions':[]}


def synthetic_chart(tmp_path):
    import os
    configured = os.environ.get('JFW_TEST_SANS_FONT')
    candidates = [configured] if configured else []
    candidates += ['C:/Windows/Fonts/arial.ttf', '/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf']
    font_path = next((path for path in candidates if path and Path(path).is_file()), None)
    if not font_path:
        pytest.skip('set JFW_TEST_SANS_FONT to an installed sans-serif font')
    f=ImageFont.truetype(font_path,24)
    im=Image.new('RGB',(800,580),'white');d=ImageDraw.Draw(im);d.line([(140,80),(140,470),(730,470)],fill='black',width=3)
    for i,x in enumerate([140,330,520,710]):
        d.line([(x,470),(x,459)],fill='black',width=3);d.text((x-15,484),str(i*1000),font=f,fill='black')
    for i,y in enumerate([80,210,340,470]):
        d.line([(140,y),(151,y)],fill='black',width=3);d.text((83,y-12),str(3-i),font=f,fill='black')
    d.line([(190,360),(330,300),(530,350),(690,390)],fill='black',width=3)
    for y,label in [(118,'Alpha'),(166,'Beta')]:
        d.line([(480,y),(540,y)],fill='black',width=3);d.text((555,y-12),label,font=f,fill='black')
    p=tmp_path/'chart.png';im.save(p);return p


def test_ranges_edges():assert ranges([0,1,1,0,1])==[(1,3),(4,5)]
@pytest.mark.parametrize('vals,expected',[([],[]),([True],[(0,1)]),([False],[]),([1,1,1],[(0,3)])])
def test_ranges(vals,expected):assert ranges(vals)==expected


def test_raster_geometry_detects_axes_not_data(tmp_path):
    g=raster_candidates(synthetic_chart(tmp_path));a=g['axes'][0]
    assert abs(a['frame'][0]-140)<=2 and abs(a['frame'][3]-470)<=2
    assert len(a['x_ticks'])>=4 and len(a['y_ticks'])>=4
    assert g['analysis']['ocr_calls']==0 and not g['analysis'].get('curve_digitization',True)
    assert len(g['legend_rows'])>=2


def test_blank_raster_returns_no_invented_axes(tmp_path):
    p=tmp_path/'blank.png';Image.new('RGB',(300,200),'white').save(p)
    g=raster_candidates(p);assert not g['axes'] and not g['legend_rows'] and g['warnings']


def test_pixel_limit(tmp_path):
    p=tmp_path/'x.png';Image.new('L',(30,30),'white').save(p)
    with pytest.raises(FigureError,match='上限'):raster_candidates(p,max_pixels=400)

@pytest.mark.parametrize('d',["M 1 2 L 3 4 Z","m 1 2 l 2 2 z"])
def test_original_linear_geometry(d):assert simple_path_points(d)==[[1.,2.],[3.,4.]]
@pytest.mark.parametrize('d',['M 0 0 Q 1 2 3 4','M 0 0 C 1 2 3 4 5 6','M 0 0 A 3 3 0 0 0 2 2'])
def test_do_not_digitize_curve(d):assert simple_path_points(d)==[]


def test_group_transform_scaling_and_offset():
    x=ET.fromstring(b'<a:xfrm xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:off x="100" y="200"/><a:ext cx="400" cy="600"/><a:chOff x="10" y="20"/><a:chExt cx="200" cy="200"/></a:xfrm>')
    m=group_matrix(x);assert point(m,10,20)==[100,200];assert point(m,210,220)==[500,800]


def test_group_rotation():
    x=ET.fromstring(b'<a:xfrm rot="5400000" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:off x="0" y="0"/><a:ext cx="100" cy="100"/><a:chOff x="0" y="0"/><a:chExt cx="100" cy="100"/></a:xfrm>')
    assert point(group_matrix(x),0,0)==pytest.approx([100,0])


def test_native_pdf_extract_not_ocr(tmp_path):
    p=source_pdf(tmp_path);dest=tmp_path/'ev';dest.mkdir();pages=extract_native(p,dest)
    texts=[q for q in pages[0]['objects'] if q['kind']=='native-text'];assert [q['text'] for q in texts]==['Input','Check','Output']
    assert pages[0]['counts']['paths']==3 and (dest/pages[0]['preview']).is_file()


def test_pptx_group_world_coordinates(tmp_path):
    prs=Presentation();s=prs.slides.add_slide(prs.slide_layouts[6]);g=s.shapes.add_group_shape()
    t=g.shapes.add_textbox(Inches(1),Inches(1),Inches(1),Inches(.5));t.text='Native grouped'
    g.left=Inches(2);g.top=Inches(3);g.width=Inches(2);g.height=Inches(1)
    p=tmp_path/'x.pptx';prs.save(p);dest=tmp_path/'ev';dest.mkdir();pages=extract_native(p,dest)
    q=next(q for q in pages[0]['objects'] if q.get('text'));assert q['rect']==pytest.approx([144,216,288,288]);assert q['parent_id']


def test_prepare_no_fonts_and_no_mutation(tmp_path):
    src=source_pdf(tmp_path);sha=digest_file(src);job=tmp_path/'job';r=prepare([src],job)
    assert r['status']=='awaiting-agent' and digest_file(src)==sha
    assert read_json(job/'job.json')['timing']['font_scans']==0
    assert not (job/'output').exists() and not list(job.rglob('*.ttf'))


def test_duplicate_source_only_extracted_once(tmp_path):
    src=source_pdf(tmp_path);r=prepare([src,src],tmp_path/'job');assert r['sources']==1


def test_prepare_never_overwrites(tmp_path):
    src=source_pdf(tmp_path);job=tmp_path/'job';prepare([src],job)
    with pytest.raises(FigureError,match='新建'):prepare([src],job)

@pytest.mark.parametrize('workers',[0,5,-1,True,1.5])
def test_worker_limits(tmp_path,workers):
    with pytest.raises(FigureError):prepare([source_pdf(tmp_path)],tmp_path/'job',workers=workers)


def test_svg_external_resource_blocked(tmp_path):
    p=tmp_path/'x.svg';p.write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10"><image href="https://example.org/steal"/></svg>')
    prepare([p],tmp_path/'job');ev=read_json(tmp_path/'job/evidence/source-1/evidence.json')
    assert ev['status']=='blocked' and ev['error']['code']=='svg_external_resource'


def test_generic_graph_measures_before_drawing(reg):
    spec={'nodes':[{'id':'a','text':'Start','shape':'roundrect'},{'id':'b','text':'Check?','shape':'diamond'},{'id':'c','text':'Done'}],
          'edges':[{'from':'a','to':'b'},{'from':'b','to':'c'}]}
    s=graph(spec,reg,8);assert not audit_scene(s);assert s.size==8
    assert s.layout_evidence['global_optimum_proven'] is False


def test_graph_cycle_not_silently_reversed(reg):
    with pytest.raises(FigureError) as exc:graph({'nodes':[{'id':'a','text':'A'},{'id':'b','text':'B'}],'edges':[{'from':'a','to':'b'},{'from':'b','to':'a'}]},reg,8)
    assert exc.value.code=='graph_cycle_requires_semantics'


def test_graph_return_edge(reg):
    s=graph({'nodes':[{'id':'a','text':'A'},{'id':'b','text':'B'}],'edges':[{'from':'a','to':'b'},{'from':'b','to':'a','return_edge':True}]},reg,8)
    assert not audit_scene(s)


def test_graph_diamond_calibration(reg):
    from journal_pdf.editorial import diamond_size
    m=reg.measure('达到要求？',8);w,h,_=diamond_size(m['width'],m['height'],8)
    s=graph({'nodes':[{'id':'d','text':'达到要求？','shape':'diamond'}],'edges':[]},reg,8)
    r=s.nodes[0]['rect'];assert [r[2]-r[0],r[3]-r[1]]==pytest.approx([w,h])


def test_conversation_roundtrip_to_actual_pdf(tmp_path,profile):
    src=source_pdf(tmp_path);job=tmp_path/'job';prepare([src],job);reply=reply_for_job(job)
    result=complete(job,reply,profile,cache=tmp_path/'cache')
    assert result['rendered']==1 and not result['failures'];r=result['results'][0]
    assert r['pdf_audit']['text_multiset_exact'] and r['layout']['layout_evidence']['layout_build_count']==1
    assert result['timings']['includes_agent_and_tool_wait'] and result['timings']['visual_review_included'] is False
    assert result['status']=='awaiting-visual-review'
    p=job/'attempt-01/output'/r['output_pdf'];assert p.is_file()


def test_one_provider_callback_is_supported(tmp_path,profile):
    src=source_pdf(tmp_path);job=tmp_path/'job';calls=[]
    def provider(request):
        assert request['is_compact_view'] is True
        calls.append(request['job_id']);return reply_for_job(job)
    result=run_with_agent([src],job,profile,provider,cache=tmp_path/'cache')
    assert len(calls)==1 and result['rendered']==1

@pytest.mark.parametrize('mutation,code',[
    (lambda r:r.update(request_sha256='wrong'),'reply_binding'),
    (lambda r:r.update(protocol='wrong'),'reply_binding'),
    (lambda r:r.update(command='delete'),'reply_schema'),
])
def test_reply_binding_security(tmp_path,profile,mutation,code):
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);reply=reply_for_job(job);mutation(reply)
    with pytest.raises(FigureError) as exc:complete(job,reply,profile)
    assert exc.value.code==code

@pytest.mark.parametrize('target',['input/source-1.pdf','evidence/source-1/evidence.json','agent-request.json'])
def test_changed_bound_evidence_fails_closed(tmp_path,profile,target):
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);reply=reply_for_job(job)
    with (job/target).open('ab') as f:f.write(b' ')
    with pytest.raises(FigureError):complete(job,reply,profile)

@pytest.mark.parametrize('case,code',[
    ('no-review','source_review_required'),('text-omitted','unaccounted_source_text'),
    ('fake-coordinate','semantic_unknown_fields'),('unbound-ref','candidate_ref_invalid'),
    ('change-native','native_text_changed'),('outside-scope','scope_ref_invalid'),
])
def test_semantic_compiler_guards(tmp_path,profile,case,code):
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);reply=reply_for_job(job);f=reply['figures'][0]
    if case=='no-review':f['review']={}
    elif case=='text-omitted':f['nodes'].pop();f['edges'].pop()
    elif case=='fake-coordinate':f['plot_clip']=[0,0,50,50]
    elif case=='unbound-ref':f['nodes'][0]['ref']='fake'
    elif case=='change-native':f['nodes'][0]['text']='Changed'
    elif case=='outside-scope':f.update(scope_refs=['fake'],scope_reason='x')
    r=complete(job,reply,profile);assert r['rendered']==0 and r['failures'][0]['error']['code']==code


def test_one_failure_does_not_block_other_figures(tmp_path,profile):
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);reply=reply_for_job(job);bad=copy.deepcopy(reply['figures'][0]);bad['id']='bad';bad['nodes'][0]['ref']='absent';reply['figures'].append(bad)
    result=complete(job,reply,profile,cache=tmp_path/'cache');assert result['rendered']==1 and len(result['failures'])==1


def test_local_retry_reuses_only_unchanged_pdf(tmp_path,profile):
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);reply=reply_for_job(job)
    first=complete(job,reply,profile,cache=tmp_path/'cache');second=complete(job,reply,profile,cache=tmp_path/'cache')
    assert first['reused_unaffected_figures']==0 and second['reused_unaffected_figures']==1
    assert second['results'][0]['pdf_sha256']==first['results'][0]['pdf_sha256']


def test_retry_limit_not_infinite(tmp_path,profile):
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);reply=reply_for_job(job);reply['figures'][0]['review']={}
    for _ in range(3):complete(job,reply,profile)
    with pytest.raises(FigureError) as exc:complete(job,reply,profile)
    assert exc.value.code=='retry_limit'


def test_visual_review_needs_real_evidence(tmp_path,profile):
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);result=complete(job,reply_for_job(job),profile,cache=tmp_path/'cache')
    with pytest.raises(FigureError):review(job,reviewer='tester',items=[{'id':'native-flow','author':'Synthetic','passed':True}])
    r=result['results'][0];ev='attempt-01/output/'+r['preview'];record=review(job,reviewer='test-only',items=[{'id':r['id'],'author':r['author'],'passed':True,'evidence':ev,'notes':'synthetic test only'}])
    assert record['ready_for_publication'] is False
    assert len(record['items'][0]['evidence_sha256'])==64 and record['elapsed_input_to_visual_review_ms']>0


def test_font_cache_no_rescan_on_new_job(tmp_path,profile):
    src=source_pdf(tmp_path);cache=tmp_path/'cache'
    for i in range(2):
        job=tmp_path/f'job{i}';prepare([src],job);r=complete(job,reply_for_job(job),profile,cache=cache)
    assert r['font_cache']['stats']['directory_scans']==0


def test_native_source_partial_scope_requires_reason(tmp_path,profile):
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);reply=reply_for_job(job);f=reply['figures'][0]
    f['scope_refs']=[n['ref'] for n in f['nodes']]
    r=complete(job,reply,profile);assert r['failures'][0]['error']['code']=='scope_reason_required'


def test_engine_version_invalidates_old_request(tmp_path,profile,monkeypatch):
    import journal_pdf.agent_entry as ae
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);reply=reply_for_job(job)
    monkeypatch.setattr(ae,'engine_signature',lambda:'different')
    with pytest.raises(FigureError) as exc:complete(job,reply,profile)
    assert exc.value.code=='stale_engine'


def test_failed_figures_do_not_create_empty_pdf(tmp_path,profile):
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);reply=reply_for_job(job);reply['figures'][0]['review']={}
    r=complete(job,reply,profile);assert not list((job/'attempt-01/output').rglob('*.pdf'))


def test_explicit_lock_is_respected(tmp_path,profile):
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);(job/'.agent.lock').write_text('foreign')
    with pytest.raises(FigureError) as exc:complete(job,reply_for_job(job),profile)
    assert exc.value.code=='job_locked' and (job/'.agent.lock').read_text()=='foreign'


def test_raster_compilation_does_not_accept_unchecked_mask(tmp_path):
    src=synthetic_chart(tmp_path);job=tmp_path/'job';prepare([src],job);state=read_json(job/'job.json');page=read_json(job/'evidence/source-1/evidence.json')['pages'][0];g=page['geometry'];a=g['axes'][0]
    f={'id':'chart','kind':'raster-chart','axis_ref':a['id'],
       'review':{'source_checked':True,'labels_checked':True,'inward_ticks_confirmed':True},
       'x_ticks':[{'ref':a['x_ticks'][0]['id'],'text':'0'}],'y_ticks':[{'ref':a['y_ticks'][0]['id'],'text':'3'}],
       'x_title':'Time(s)','y_title':'Power(W)','legend_rows':[{'ref':g['legend_rows'][0]['id'],'label':'Alpha','series_id':'a'}]}
    with pytest.raises(FigureError) as exc:compile_reply_figure(f,state['sources'][0],page)
    assert exc.value.code=='mask_semantics_required'


def test_raster_no_silent_outward_tick_clip(tmp_path):
    src=synthetic_chart(tmp_path);job=tmp_path/'job';prepare([src],job);state=read_json(job/'job.json');page=read_json(job/'evidence/source-1/evidence.json')['pages'][0]
    f={'id':'chart','kind':'raster-chart','review':{'source_checked':True,'labels_checked':True}}
    with pytest.raises(FigureError) as exc:compile_reply_figure(f,state['sources'][0],page)
    assert exc.value.code=='raster_tick_direction_pending'


def test_native_rule_ledger_keeps_preview_pending(tmp_path,profile):
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);r=complete(job,reply_for_job(job),profile,cache=tmp_path/'cache')
    ledger=r['results'][0]['rule_evidence'];states={x['id']:x['status'] for x in ledger['rules']}
    assert states['production-composite-font-equivalence']=='pending';assert states['whole-artwork-default8-min6']=='checked'
    assert states['axis-quantity-slash-unit']=='not-applicable'


def test_structured_tree_uses_text_refs_only(tmp_path,profile):
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);r=reply_for_job(job);f=r['figures'][0];refs=[n['ref'] for n in f.pop('nodes')];f.pop('edges');f['model']='tree'
    f['structure']={'root':{'ref':refs[0]},'branches':[{'text':{'ref':refs[1]},'children':[{'ref':refs[2]}]}]}
    result=complete(job,r,profile,cache=tmp_path/'cache');assert result['rendered']==1 and not result['failures']


def test_native_shape_source_cannot_be_hijacked(tmp_path,profile):
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);r=reply_for_job(job);r['figures'][0]['nodes'][0]['ref']='p1-path-0'
    result=complete(job,r,profile);assert result['failures'][0]['error']['code']=='native_text_missing'


def test_compact_request_omits_precise_geometry(tmp_path):
    job=tmp_path/'job';prepare([synthetic_chart(tmp_path)],job)
    brief=read_json(job/'agent-brief.json');text=json.dumps(brief)
    assert 'plot_clip' not in text and 'swatch_clip' not in text and 'is_compact_view' in text
    assert brief['request_sha256']==read_json(job/'agent-request.json')['request_sha256']


def test_modified_compact_brief_is_rejected(tmp_path,profile):
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);reply=reply_for_job(job)
    with (job/'agent-brief.json').open('a') as f:f.write(' ')
    with pytest.raises(FigureError) as exc:complete(job,reply,profile)
    assert exc.value.code=='brief_changed'


def test_multi_frame_input_not_silently_truncated(tmp_path):
    p=tmp_path/'multi.tiff';a=Image.new('RGB',(100,100),'white');a.save(p,save_all=True,append_images=[a])
    with pytest.raises(FigureError) as exc:raster_candidates(p)
    assert exc.value.code=='raster_multiframe'


def test_rotated_pdf_keeps_coordinate_contract(tmp_path):
    src=tmp_path/'rot.pdf';d=fitz.open();p=d.new_page(width=200,height=100);p.insert_text((20,40),'Hello');p.set_rotation(90);d.save(src);d.close()
    folder=tmp_path/'ev';folder.mkdir();page=extract_native(src,folder)[0]
    assert page['size']==[200,100] and page['display_size']==[100,200] and page['rotation']==90


def test_pdf_identity_cannot_escape_source_root(tmp_path,profile):
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);reply=reply_for_job(job);reply['figures'][0]['source']='../../unrelated.pdf'
    r=complete(job,reply,profile);assert r['failures'][0]['error']['code']=='semantic_unknown_fields'


def test_actual_font_identity_invalidates_partial_reuse(tmp_path,profile,monkeypatch):
    import journal_pdf.fonts as fm
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);rep=reply_for_job(job);complete(job,rep,profile,cache=tmp_path/'cache')
    original=fm.FontRegistry
    class ChangedIdentity(original):
        def __init__(self,*args,**kwargs):
            super().__init__(*args,**kwargs);self.fingerprint+='test-changed-font-identity'
    monkeypatch.setattr(fm,'FontRegistry',ChangedIdentity)
    result=complete(job,rep,profile,cache=tmp_path/'cache');assert result['rendered']==1 and result['reused_unaffected_figures']==0


def test_changed_previous_preview_is_not_reused(tmp_path,profile):
    job=tmp_path/'job';prepare([source_pdf(tmp_path)],job);rep=reply_for_job(job);r=complete(job,rep,profile,cache=tmp_path/'cache')
    preview=job/'attempt-01/output'/r['results'][0]['preview']
    with preview.open('ab') as f:f.write(b'changed')
    second=complete(job,rep,profile,cache=tmp_path/'cache')
    assert second['failures'][0]['error']['code']=='reuse_preview_changed'
